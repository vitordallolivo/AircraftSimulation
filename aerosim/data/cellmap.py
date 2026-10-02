"""Planilhas "por endereço de célula" (ex.: Dados_Desempenho.xlsx), incluindo os conjuntos de
tração por polinômio da aba Tracao.

Em vez de código com `sheet["C4"].value` espalhado, um ARQUIVO DE MAPA (.json) diz qual
campo do modelo vem de qual célula, com escala. Trocar a planilha = trocar o mapa.

    {
      "base": "conventional",                       # aeronave de partida (template ou arquivo)
      "fields": {
        "mass.base_mass": "Geral!C8",
        "surfaces[wing].CLa": {"cell": "Superficies!C4", "scale": 57.29577951308232,
                               "note": "planilha em 1/grau"},
        "@wing.area": "Superficies!C3",              # geometria derivada (refaz as seções)
        ...
      }
    }

Caminhos: `a.b.c`, listas por índice `propulsion[0]` ou por papel `surfaces[wing]`,
vetores `mass.base_cg[0]`. Chaves com `@papel.` ajustam a geometria da superfície:
area, span, xac (centro aerodinâmico), chord, K (fator de arrasto induzido -> Oswald).
"set": {"tanks": [], ...} aplica valores constantes antes de ler as células.

Conjuntos de tração (uma linha por conjunto motor+hélice, coeficientes p00..p05 do poly55):

    "thrust_sets": {"sheet": "Tracao", "first_row": 3, "name_col": 1, "coef_first_col": 2,
                    "n_col": 23, "T0_col": 25, "select_cell": "Geral!C28", "row_offset": 2,
                    "x": "V", "y": "none", "gain": "1"}

Todas as linhas viram o catálogo `self.catalog.thrust_sets`; a linha `select_cell + row_offset`
(o `nhelice` da planilha) vai para todas as unidades propulsoras da aeronave.
"""
from __future__ import annotations

import json
import math
import os
import re

from aerosim.interfaces.aircraft_source import IAircraftSource, ICatalogSource
from aerosim.model.aircraft import Aircraft, LiftingSurface, WingSection
from aerosim.model.catalog import Catalog, ThrustPoly

_TOK = re.compile(r"([A-Za-z_][A-Za-z_0-9]*)(?:\[([^\]]+)\])?")


def _split_cell(ref: str) -> tuple[str, str]:
    sheet, _, cell = ref.rpartition("!")
    return sheet, cell


def _resolve(root, path: str):
    """devolve (objeto pai, atributo ou índice) do caminho"""
    toks = path.split(".")
    obj = root
    for i, t in enumerate(toks):
        m = _TOK.fullmatch(t)
        if not m:
            raise KeyError(f"caminho inválido: {path}")
        name, sel = m.group(1), m.group(2)
        last = i == len(toks) - 1
        if sel is None:
            if last:
                return obj, name
            obj = getattr(obj, name)
            continue
        seq = getattr(obj, name)
        if sel.lstrip("-").isdigit():
            idx = int(sel)
        else:
            idx = next((k for k, it in enumerate(seq) if getattr(it, "role", None) == sel or getattr(it, "name", None) == sel), None)
            if idx is None:
                raise KeyError(f"{path}: nenhum item '{sel}' em {name}")
        if last:
            return seq, idx
        obj = seq[idx]
    raise KeyError(path)


def _set(root, path, value):
    parent, key = _resolve(root, path)
    if isinstance(parent, list):
        parent[key] = value
    else:
        cur = getattr(parent, key)
        if isinstance(cur, bool):
            value = bool(value)
        elif isinstance(cur, int) and not isinstance(cur, bool):
            value = int(value)
        setattr(parent, key, value)


def _get(root, path):
    parent, key = _resolve(root, path)
    return parent[key] if isinstance(parent, list) else getattr(parent, key)


def _surface(ac: Aircraft, role: str) -> LiftingSurface:
    s = next((s for s in ac.surfaces if s.role == role or s.name == role), None)
    if s is None:
        raise KeyError(f"superfície '{role}' não existe na aeronave base")
    return s


def _reshape(s: LiftingSurface, area: float | None = None, span: float | None = None, chord: float | None = None):
    """refaz as seções mantendo afilamento, enflechamento e torção"""
    ss = sorted(s.sections, key=lambda q: q.y)
    root, tip = ss[0], ss[-1]
    taper = tip.chord / root.chord if root.chord > 0 else 1.0
    mult = 1 if (s.vertical or not s.symmetric) else 2
    b = span if span is not None else s.span
    semi = b / mult
    if area is not None:
        cmean = area / b
    elif chord is not None:
        cmean = chord * (1 + taper) / 2
    else:
        cmean = s.area / b if b > 0 else root.chord
    c_root = 2 * cmean / (1 + taper)
    sweep_per_y = (tip.x_offset - root.x_offset) / tip.y if tip.y > 0 else 0.0
    s.sections = [WingSection(f * semi, c_root * (1 + (taper - 1) * f), f * semi * sweep_per_y,
                              root.twist + (tip.twist - root.twist) * f) for f in (0.0, 1.0)]


class CellMapAircraftSource(IAircraftSource):
    extensions = (".xlsx",)

    def __init__(self, mapping: str | dict, base: Aircraft | None = None):
        if isinstance(mapping, str):
            with open(mapping, encoding="utf-8") as f:
                mapping = json.load(f)
        self.mapping = mapping
        self.base = base
        self.warnings: list[str] = []
        self.catalog = Catalog()          # conjuntos de tração lidos da planilha
        self._sel_row: int | None = None

    # ------------------------------------------------------------------ conjuntos de tração
    @staticmethod
    def _cell_value(wb, ref):
        sheet, cell = _split_cell(ref)
        return wb[sheet][cell].value if sheet in wb.sheetnames else None

    def _selected_row(self, wb, ts) -> int | None:
        if "select_cell" not in ts:
            return None
        v = self._cell_value(wb, ts["select_cell"])
        return int(v) + int(ts.get("row_offset", 0)) if isinstance(v, (int, float)) else None

    def _gain(self, wb, ts, T0) -> float:
        g = str(ts.get("gain", "1"))
        if g == "T0":
            return T0
        if g == "T0*PT0":
            pt0 = self._cell_value(wb, ts.get("pt0_cell", "Geral!C29"))
            return T0 * (float(pt0) if isinstance(pt0, (int, float)) else 1.0)
        return float(g)

    def _write_thrust_sets(self, wb, ts, aircraft: Aircraft) -> None:
        """grava os conjuntos (o da aeronave na linha selecionada) e a célula de seleção"""
        sel = self._selected_row(wb, ts) or self._sel_row
        u = next((x for x in aircraft.propulsion if x.thrust_model == "poly"), None)
        if u is None:
            return
        first = int(ts.get("first_row", 3))
        if sel is None:
            sel = first
        if ts["sheet"] not in wb.sheetnames:
            wb.create_sheet(ts["sheet"])
        sh = wb[ts["sheet"]]
        sets = {first + i: t for i, t in enumerate(self.catalog.thrust_sets)}
        sets[sel] = u.thrust_poly
        c0 = int(ts.get("coef_first_col", 2))
        for row, t in sets.items():
            sh.cell(row=row, column=int(ts.get("name_col", 1)), value=t.name)
            for k, c in enumerate((list(t.coefs) + [0.0] * 21)[:21]):
                sh.cell(row=row, column=c0 + k, value=c)
            if "n_col" in ts:
                sh.cell(row=row, column=int(ts["n_col"]), value=t.n)
            if "T0_col" in ts:
                sh.cell(row=row, column=int(ts["T0_col"]), value=t.T0)
        if "select_cell" in ts:
            sheet, cell = _split_cell(ts["select_cell"])
            if sheet not in wb.sheetnames:
                wb.create_sheet(sheet)
            wb[sheet][cell] = sel - int(ts.get("row_offset", 0))

    def _read_thrust_sets(self, wb, ac: Aircraft) -> None:
        ts = self.mapping.get("thrust_sets")
        if not ts:
            return
        if ts["sheet"] not in wb.sheetnames:
            self.warnings.append(f"thrust_sets: aba '{ts['sheet']}' não existe")
            return
        sh = wb[ts["sheet"]]
        c0, ncoef = int(ts.get("coef_first_col", 2)), 21
        rows = {}
        r = int(ts.get("first_row", 3))
        while True:
            coefs = [sh.cell(row=r, column=c0 + k).value for k in range(ncoef)]
            if all(v is None for v in coefs):
                break
            name = sh.cell(row=r, column=int(ts.get("name_col", 1))).value
            T0 = sh.cell(row=r, column=int(ts["T0_col"])).value if "T0_col" in ts else None
            n = sh.cell(row=r, column=int(ts["n_col"])).value if "n_col" in ts else None
            T0 = float(T0) if isinstance(T0, (int, float)) else 0.0
            tp = ThrustPoly(name=str(name) if name not in (None, "") else f"conjunto {r - int(ts.get('row_offset', 0))}",
                            coefs=[float(v) if isinstance(v, (int, float)) else 0.0 for v in coefs],
                            x_var=ts.get("x", "V"), y_var=ts.get("y", "none"), gain=self._gain(wb, ts, T0),
                            rho_exp=float(ts.get("rho_exp", 0.0)), xmax=float(ts.get("xmax", 0.0)),
                            ymax=float(ts.get("ymax", 0.0)), T0=T0, n=float(n) if isinstance(n, (int, float)) else 0.0,
                            source=f"{ts['sheet']}!linha {r}")
            rows[r] = tp
            self.catalog.thrust_sets.append(tp)
            r += 1
        sel = self._selected_row(wb, ts)
        self._sel_row = sel
        if sel is None:
            return
        if sel not in rows:
            self.warnings.append(f"thrust_sets: linha {sel} (nhelice) não tem coeficientes")
            return
        import copy as _copy
        for u in ac.propulsion:
            u.thrust_model = "poly"
            u.thrust_poly = _copy.deepcopy(rows[sel])
        if not ac.propulsion:
            self.warnings.append("thrust_sets: a aeronave base não tem motor para receber o conjunto")

    def _base(self) -> Aircraft:
        if self.base is not None:
            return self.base.copy()
        b = self.mapping.get("base", "conventional")
        from aerosim.model.templates import TEMPLATES
        if b in TEMPLATES:
            return TEMPLATES[b]()
        from .sources import load_aircraft
        return load_aircraft(b)

    @staticmethod
    def _spec(v) -> dict:
        return {"cell": v} if isinstance(v, str) else dict(v)

    def load(self, path: str) -> Aircraft:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ac = self._base()
        self.warnings = []
        geo: dict[str, dict] = {}
        for target, value in self.mapping.get("set", {}).items():      # constantes (ex.: tirar tanques)
            try:
                _set(ac, target, value)
            except (KeyError, IndexError, AttributeError) as e:
                self.warnings.append(f"set {target}: {e}")
        for target, raw in self.mapping.get("fields", {}).items():
            sp = self._spec(raw)
            sheet, cell = _split_cell(sp["cell"])
            if sheet not in wb.sheetnames:
                self.warnings.append(f"{target}: aba '{sheet}' não existe")
                continue
            v = wb[sheet][cell].value
            if v is None or isinstance(v, str) and not v.strip():
                self.warnings.append(f"{target}: {sp['cell']} vazia")
                continue
            if isinstance(v, (int, float)):
                v = float(v) * sp.get("scale", 1.0) + sp.get("offset", 0.0)
            try:
                if target.startswith("@"):
                    role, _, what = target[1:].partition(".")
                    geo.setdefault(role, {})[what] = v
                else:
                    _set(ac, target, v)
            except (KeyError, IndexError, AttributeError, TypeError, ValueError) as e:
                self.warnings.append(f"{target}: {e}")
        for role, g in geo.items():
            try:
                s = _surface(ac, role)
            except KeyError as e:
                self.warnings.append(str(e))
                continue
            if any(k in g for k in ("area", "span", "chord")):
                _reshape(s, g.get("area"), g.get("span"), g.get("chord"))
            if "xac" in g:
                s.x = g["xac"] + 0.25 * s.mac
            if "K" in g and g["K"] > 0:          # C_Di = K C_L²  ->  e = 1 / (π AR K)
                s.oswald = 1.0 / (math.pi * s.aspect_ratio * g["K"])
        self.catalog = Catalog()
        self._read_thrust_sets(wb, ac)
        wb.close()
        ac.apply_constraints()
        return ac

    def save(self, aircraft: Aircraft, path: str) -> None:
        """escreve os valores de volta nas mesmas células (abre a planilha se ela existir)"""
        from openpyxl import Workbook, load_workbook
        wb = load_workbook(path) if os.path.exists(path) else Workbook()
        for target, raw in self.mapping.get("fields", {}).items():
            sp = self._spec(raw)
            sheet, cell = _split_cell(sp["cell"])
            if sheet not in wb.sheetnames:
                wb.create_sheet(sheet)
            try:
                if target.startswith("@"):
                    role, _, what = target[1:].partition(".")
                    s = _surface(aircraft, role)
                    v = {"area": s.area, "span": s.span, "chord": s.mac, "xac": s.x - 0.25 * s.mac,
                         "K": 1.0 / (math.pi * s.aspect_ratio * s.oswald)}[what]
                else:
                    v = _get(aircraft, target)
            except (KeyError, IndexError, AttributeError):
                continue
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                v = (v - sp.get("offset", 0.0)) / sp.get("scale", 1.0)
            if v is None:
                continue
            wb[sheet][cell] = v
        ts = self.mapping.get("thrust_sets")
        if ts:
            self._write_thrust_sets(wb, ts, aircraft)
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1 and wb["Sheet"].max_row == 1:
            del wb["Sheet"]
        wb.save(path)
