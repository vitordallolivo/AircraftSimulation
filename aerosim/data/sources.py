"""Implementações de IAircraftSource e ICatalogSource: .json, .cfg (INI) e .xlsx (Excel).

    load_aircraft(path) / save_aircraft(ac, path)     escolhem o formato pela extensão
    load_catalog(path)  / save_catalog(cat, path)
"""
from __future__ import annotations

import configparser
import json
import os
import re
from collections import defaultdict

from aerosim.interfaces.aircraft_source import IAircraftSource, ICatalogSource
from aerosim.model.aircraft import Aircraft
from aerosim.model.catalog import POLY_TERMS, Battery, Catalog, Motor, Propeller, ThrustPoly
from aerosim.model.serialize import from_dict, to_dict

from .flat import decode, encode, flatten, unflatten


def _strip_coords(d):
    if isinstance(d, dict):
        return {k: _strip_coords(v) for k, v in d.items() if k != "coords"}
    if isinstance(d, list):
        return [_strip_coords(v) for v in d]
    return d


def _resolve_paths(ac: Aircraft, base: str) -> Aircraft:
    """caminhos de .dat relativos ao arquivo da aeronave"""
    for s in ac.surfaces:
        p = s.airfoil.path
        if p and not os.path.isabs(p):
            cand = os.path.normpath(os.path.join(base, p))
            if os.path.exists(cand):
                s.airfoil.path = cand
    return ac


# ================================================================== aeronave
class JsonAircraftSource(IAircraftSource):
    extensions = (".json",)

    def load(self, path):
        with open(path, encoding="utf-8") as f:
            return _resolve_paths(from_dict(Aircraft, json.load(f)), os.path.dirname(path))

    def save(self, aircraft, path):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(_strip_coords(to_dict(aircraft)), f, indent=1, ensure_ascii=False)


class CfgAircraftSource(IAircraftSource):
    """INI legível: uma seção por componente, listas como [a, b, c]"""
    extensions = (".cfg", ".ini")

    def load(self, path):
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        cp.read(path, encoding="utf-8")
        sections = {s: {k: decode(v) for k, v in cp[s].items()} for s in cp.sections()}
        return _resolve_paths(from_dict(Aircraft, unflatten(sections)), os.path.dirname(path))

    def save(self, aircraft, path):
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        for sec, keys in flatten(aircraft).items():
            cp[sec] = {k: encode(v) for k, v in keys.items()}
        with open(path, "w", encoding="utf-8") as f:
            f.write("# aeronave (aerosim) — unidades SI, ângulos em graus; listas em [ ]\n")
            cp.write(f)


_GROUP = re.compile(r"\.(\d+)")


def _group_name(section: str) -> tuple[str, list[int]]:
    """'surfaces.0.sections.1' -> ('surfaces.#.sections.#', [0, 1])"""
    idx = [int(m) for m in _GROUP.findall(section)]
    return _GROUP.sub(".#", section), idx


class ExcelAircraftSource(IAircraftSource):
    """planilha: aba 'geral' (seção, chave, valor) + uma aba-tabela por lista de componentes"""
    extensions = (".xlsx",)

    def save(self, aircraft, path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "geral"
        ws.append(["seção", "chave", "valor"])
        groups: dict[str, list] = defaultdict(list)
        for sec, keys in flatten(aircraft).items():
            g, idx = _group_name(sec)
            if not idx:
                for k, v in keys.items():
                    ws.append([sec, k, self._cell(v)])
            else:
                groups[g].append((idx, keys))
        for g, rows in groups.items():
            sh = wb.create_sheet(g[:31])
            depth = len(rows[0][0])
            cols = []
            for _, keys in rows:
                cols += [k for k in keys if k not in cols]
            sh.append([f"idx{i}" for i in range(depth)] + cols)
            for idx, keys in rows:
                sh.append(idx + [self._cell(keys.get(c)) if c in keys else None for c in cols])
        wb.save(path)

    @staticmethod
    def _cell(v):
        """números ficam números na planilha; o resto vai como texto (JSON quando preciso)"""
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return v
        return encode(v)

    def load(self, path):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        sections: dict[str, dict] = defaultdict(dict)
        val = lambda v: decode(v) if isinstance(v, str) else v
        for row in wb["geral"].iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            sections[row[0]][row[1]] = val(row[2]) if row[2] is not None else ""
        for sh in wb.worksheets:
            if sh.title == "geral":
                continue
            rows = list(sh.iter_rows(values_only=True))
            if not rows:
                continue
            head = rows[0]
            depth = sum(1 for h in head if str(h).startswith("idx"))
            for r in rows[1:]:
                if r[0] is None:
                    continue
                name = sh.title
                for v in r[:depth]:
                    name = name.replace("#", str(int(v)), 1)
                sections[name] = {h: val(v) for h, v in zip(head[depth:], r[depth:]) if v is not None}
        return _resolve_paths(from_dict(Aircraft, unflatten(dict(sections))), os.path.dirname(path))


# ================================================================== catálogo
class JsonCatalogSource(ICatalogSource):
    extensions = (".json",)

    def load_catalog(self, path):
        with open(path, encoding="utf-8") as f:
            return from_dict(Catalog, json.load(f))

    def save_catalog(self, catalog, path):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(to_dict(catalog), f, indent=1, ensure_ascii=False)


class CfgCatalogSource(ICatalogSource):
    """[motor: nome] / [helice: nome] (J, CT, CP em listas) / [bateria: nome] / [tracao: nome] (coefs)"""
    extensions = (".cfg", ".ini")
    KINDS = (("motor", Motor, "motors"), ("helice", Propeller, "propellers"), ("bateria", Battery, "batteries"),
             ("tracao", ThrustPoly, "thrust_sets"))

    def load_catalog(self, path):
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        cp.read(path, encoding="utf-8")
        cat = Catalog()
        for sec in cp.sections():
            kind, _, name = sec.partition(":")
            for k, cls, attr in self.KINDS:
                if kind.strip() == k:
                    d = {kk: decode(v) for kk, v in cp[sec].items()}
                    d["name"] = name.strip()
                    getattr(cat, attr).append(from_dict(cls, d))
        return cat

    def save_catalog(self, catalog, path):
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        for k, _, attr in self.KINDS:
            for item in getattr(catalog, attr):
                d = to_dict(item)
                name = d.pop("name")
                cp[f"{k}: {name}"] = {kk: encode(v) for kk, v in d.items()}
        with open(path, "w", encoding="utf-8") as f:
            f.write("# catálogo de motores, hélices e baterias (SI)\n")
            cp.write(f)


class ExcelCatalogSource(ICatalogSource):
    """abas: motores, helices (dados gerais), helices_tabela (helice, J, CT, CP), baterias,
    tracao (um conjunto por linha: nome, variáveis, ganho..., p00..p05)"""
    extensions = (".xlsx",)
    SIMPLE = (("motores", Motor, "motors"), ("baterias", Battery, "batteries"))

    def save_catalog(self, catalog, path):
        import dataclasses as dc
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        for title, cls, attr in self.SIMPLE:
            sh = wb.create_sheet(title)
            cols = [f.name for f in dc.fields(cls)]
            sh.append(cols)
            for it in getattr(catalog, attr):
                sh.append([getattr(it, c) for c in cols])
        sh = wb.create_sheet("helices")
        cols = [f.name for f in dc.fields(Propeller) if f.name not in ("J", "CT", "CP")]
        sh.append(cols)
        for p in catalog.propellers:
            sh.append([getattr(p, c) for c in cols])
        sh = wb.create_sheet("helices_tabela")
        sh.append(["helice", "J", "CT", "CP"])
        for p in catalog.propellers:
            for j, ct, cpw in zip(p.J, p.CT, p.CP):
                sh.append([p.name, j, ct, cpw])
        sh = wb.create_sheet("tracao")
        cols = [f.name for f in dc.fields(ThrustPoly) if f.name != "coefs"]
        sh.append(cols + POLY_TERMS)
        for t in catalog.thrust_sets:
            sh.append([getattr(t, c) for c in cols] + list(t.coefs))
        wb.save(path)

    def load_catalog(self, path):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)

        def table(title):
            if title not in wb.sheetnames:
                return []
            rows = list(wb[title].iter_rows(values_only=True))
            return [dict(zip(rows[0], r)) for r in rows[1:] if r and r[0] is not None]
        cat = Catalog()
        for title, cls, attr in self.SIMPLE:
            getattr(cat, attr).extend(from_dict(cls, {k: v for k, v in r.items() if v is not None}) for r in table(title))
        tab = defaultdict(list)
        for r in table("helices_tabela"):
            tab[r["helice"]].append((float(r["J"]), float(r["CT"]), float(r["CP"])))
        for r in table("helices"):
            p = from_dict(Propeller, {k: v for k, v in r.items() if v is not None})
            pts = sorted(tab.get(p.name, []))
            p.J, p.CT, p.CP = [a[0] for a in pts], [a[1] for a in pts], [a[2] for a in pts]
            cat.propellers.append(p)
        for r in table("tracao"):
            d = {k: v for k, v in r.items() if v is not None and k not in POLY_TERMS}
            d["coefs"] = [float(r.get(t) or 0.0) for t in POLY_TERMS]
            cat.thrust_sets.append(from_dict(ThrustPoly, d))
        return cat


# ================================================================== escolha pela extensão
AIRCRAFT_SOURCES: list[IAircraftSource] = [JsonAircraftSource(), CfgAircraftSource(), ExcelAircraftSource()]
CATALOG_SOURCES: list[ICatalogSource] = [JsonCatalogSource(), CfgCatalogSource(), ExcelCatalogSource()]


def _pick(sources, path):
    ext = os.path.splitext(path)[1].lower()
    for s in sources:
        if ext in s.extensions:
            return s
    raise ValueError(f"formato não suportado: {ext}")


def aircraft_source(path: str) -> IAircraftSource:
    return _pick(AIRCRAFT_SOURCES, path)


def catalog_source(path: str) -> ICatalogSource:
    return _pick(CATALOG_SOURCES, path)


def load_aircraft(path: str) -> Aircraft:
    return aircraft_source(path).load(path)


def save_aircraft(ac: Aircraft, path: str) -> None:
    aircraft_source(path).save(ac, path)


def load_catalog(path: str) -> Catalog:
    return catalog_source(path).load_catalog(path)


def save_catalog(cat: Catalog, path: str) -> None:
    catalog_source(path).save_catalog(cat, path)
