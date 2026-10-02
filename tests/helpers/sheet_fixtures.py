"""Planilha no formato do Dados_Desempenho.xlsx (abas Geral, Superficies e Tracao) para os testes."""

GERAL = {"C8": 7.5, "C5": 15, "C9": 20, "C13": 2.0, "C14": 3.0, "C18": 1.225, "C19": 1.2, "C21": 0.05, "C28": 2,
         "C29": 1.0, "G9": 0.8, "H10": 1.1, "I11": 1.7, "I9": -0.03, "G7": 0.01, "G6": 0.5, "G4": 0, "G5": 0,
         "G8": -0.06, "G12": -0.04, "G13": 0.4}
SUPERFICIES = {"M3": 0.02, "C14": 0.9, "C15": 2.4, "C16": 0.375, "C3": 0.9, "C12": 0.04, "C17": 0.06, "C4": 0.08,
               "C5": 0.35, "C6": -0.07, "C9": 2, "C13": 1.5, "C11": 1.5, "C10": 0.4, "G3": 0.18, "G4": 0.07,
               "G5": 0, "G6": 0, "G9": -1.5, "G11": 1.0, "C8": 0.35, "G8": 0.35, "J7": 0.4,
               "G36": 0.0005, "G37": 0.02, "J36": 0.0004, "J37": 0.015}
# conjuntos (polinômio em kgf, x = n, y = V):  T = g (ρ/ρ0) (T0 + p10 n + p01 V + p02 V²)
# a coluna p00 da planilha recebe lixo de propósito: o original usa T0 no lugar dela
SETS = [  # nome, T0, n, p10, p01, p02
    ("conjunto A", 2.6, 1.5, 0.10, -0.040, -0.0040),
    ("conjunto B", 2.2, 1.5, 0.10, -0.060, -0.0020),
    ("conjunto C", 2.0, 1.2, 0.05, -0.030, -0.0030),
]
TERMS = ["p00", "p10", "p01", "p20", "p11", "p02", "p30", "p21", "p12", "p03", "p40", "p31", "p22", "p13", "p04",
         "p50", "p41", "p32", "p23", "p14", "p05"]


def thrust(name, V, rho=1.2, rho0=1.225):
    """a conta do Envelope_main original (com p00 trocado por T0)"""
    _, T0, n, p10, p01, p02 = next(s for s in SETS if s[0] == name)
    return 9.80665 * (rho / rho0) * (T0 + p10 * n + p01 * V + p02 * V * V)


def write_desempenho(path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Geral"
    wb.create_sheet("Superficies")
    tr = wb.create_sheet("Tracao")
    for sh, d in (("Geral", GERAL), ("Superficies", SUPERFICIES)):
        for c, v in d.items():
            wb[sh][c] = v
    tr.cell(row=2, column=1, value="nome")
    for k, t in enumerate(TERMS):
        tr.cell(row=2, column=2 + k, value=t)
    tr.cell(row=2, column=23, value="n")
    tr.cell(row=2, column=25, value="T0")
    for i, (name, T0, n, p10, p01, p02) in enumerate(SETS):
        r = 3 + i
        tr.cell(row=r, column=1, value=name)
        for k in range(21):
            tr.cell(row=r, column=2 + k, value=0.0)
        tr.cell(row=r, column=2, value=999.0)                  # p00 (ignorado: entra T0)
        tr.cell(row=r, column=2 + TERMS.index("p10"), value=p10)
        tr.cell(row=r, column=2 + TERMS.index("p01"), value=p01)
        tr.cell(row=r, column=2 + TERMS.index("p02"), value=p02)
        tr.cell(row=r, column=23, value=n)
        tr.cell(row=r, column=25, value=T0)
    wb.save(path)
