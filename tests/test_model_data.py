"""Aircraft Data: modelo, geometria, fontes de dados (.json, .cfg, .xlsx, mapa de células)."""
import json
import math

import numpy as np
import pytest

from aerosim.data import load_aircraft, load_catalog, save_aircraft, save_catalog
from aerosim.model import TEMPLATES, builtin_catalog, compile_aircraft
from aerosim.model.airfoil import naca4, read_dat, thin_airfoil, write_dat
from aerosim.model.geometry import aircraft_meshes, all_panels, component_position, move_component


def _flat(d, p=""):
    if isinstance(d, dict):
        for k, v in d.items():
            yield from _flat(v, f"{p}.{k}")
    elif isinstance(d, list):
        for i, v in enumerate(d):
            yield from _flat(v, f"{p}[{i}]")
    else:
        yield p, d


def same_spec(a, b):
    A, B = dict(_flat(a)), dict(_flat(b))
    assert A.keys() == B.keys()
    for k in A:
        assert np.isclose(A[k], B[k], rtol=1e-12, atol=1e-14), k


@pytest.mark.parametrize("name", list(TEMPLATES))
def test_panel_area_matches_surfaces(name):
    ac = TEMPLATES[name]()
    panels = all_panels(ac)
    for g, role in ((0, "wing"), (1, "htail"), (2, "vtail")):
        S_p = sum(p["S"] for p in panels if p["group"] == g)
        S_s = sum(s.area for s in ac.surfaces if s.role == role and s.enabled)
        assert S_p == pytest.approx(S_s, rel=1e-9)


@pytest.mark.parametrize("name", list(TEMPLATES))
def test_control_and_stability_signs(backend, name):
    """mesma convenção de sinais para leme, cauda em V, elevon...; estável em guinada"""
    with backend.open_session() as s:
        s.configure(compile_aircraft(TEMPLATES[name]()))
        c0 = s.aero_coeffs(14, 0.05)
        f = lambda **kw: s.aero_coeffs(14, 0.05, **kw)
        assert f(da=0.1)["Cl"] - c0["Cl"] > 0
        assert f(de=0.1)["Cm"] - c0["Cm"] < 0
        assert f(beta=0.1)["Cn"] - c0["Cn"] > 0
        if name != "flying_wing":
            assert f(dr=0.1)["Cn"] - c0["Cn"] < 0


def test_thin_airfoil():
    a0, cm = thin_airfoil(naca4("2412"))
    assert a0 == pytest.approx(-2.08, abs=0.15)     # valor clássico da teoria de perfil fino
    assert cm == pytest.approx(-0.053, abs=0.01)
    a0s, cms = thin_airfoil(naca4("0012"))
    assert abs(a0s) < 1e-6 and abs(cms) < 1e-6


def test_dat_roundtrip(tmp_path):
    c = naca4("4412")
    p = str(tmp_path / "a.dat")
    write_dat(p, "NACA 4412", c)
    name, d = read_dat(p)
    assert name == "NACA 4412" and np.allclose(c, d, atol=1e-6)


@pytest.mark.parametrize("ext", ["json", "cfg", "xlsx"])
@pytest.mark.parametrize("name", list(TEMPLATES))
def test_aircraft_roundtrip(tmp_path, name, ext):
    ac = TEMPLATES[name]()
    p = str(tmp_path / f"a.{ext}")
    save_aircraft(ac, p)
    same_spec(compile_aircraft(ac), compile_aircraft(load_aircraft(p)))


@pytest.mark.parametrize("ext", ["json", "cfg", "xlsx"])
def test_catalog_roundtrip(tmp_path, ext):
    c = builtin_catalog()
    p = str(tmp_path / f"c.{ext}")
    save_catalog(c, p)
    d = load_catalog(p)
    assert [m.name for m in d.motors] == [m.name for m in c.motors]
    for a, b in zip(c.propellers, d.propellers):
        assert np.allclose(a.CT, b.CT) and np.allclose(a.J, b.J)


def test_move_component_and_tail_on_boom():
    ac = TEMPLATES["twin_boom"]()
    move_component(ac, "prop:1", [0.2, 0.7, -0.05])
    assert list(component_position(ac, "prop:1")) == pytest.approx([0.2, 0.7, -0.05])
    ac.boom.length += 0.2
    ac.apply_constraints()
    vt = ac.surface("htail")
    assert vt.x == pytest.approx(ac.boom.x_end + vt.sections[0].chord)
    assert "boom" in aircraft_meshes(ac)


def _desempenho_sheet(path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Geral"
    wb.create_sheet("Superficies")
    vals = {"Geral": {"C8": 7.5, "C5": 15, "C9": 20, "C19": 1.2, "C21": 0.05, "G9": 0.8, "H10": 1.1, "I11": 1.7,
                      "I9": -0.03, "G7": 0.01, "G6": 0.5, "G4": 0, "G5": 0, "G8": -0.06, "G12": -0.04, "G13": 0.4},
            "Superficies": {"M3": 0.02, "C14": 0.9, "C15": 2.4, "C16": 0.375, "C3": 0.9, "C12": 0.04, "C17": 0.06,
                            "C4": 0.08, "C5": 0.35, "C6": -0.07, "C9": 2, "C13": 1.5, "C11": 1.5, "C10": 0.4,
                            "G3": 0.18, "G4": 0.07, "G5": 0, "G6": 0, "G9": -1.5, "G11": 1.0, "C8": 0.35, "G8": 0.35,
                            "J7": 0.4}}
    for sh, d in vals.items():
        for c, v in d.items():
            wb[sh][c] = v
    wb.save(path)


def test_cellmap_source(tmp_path, manager):
    import os
    from aerosim.data.cellmap import CellMapAircraftSource
    root = os.path.join(os.path.dirname(__file__), "..")
    xl = str(tmp_path / "Dados_Desempenho.xlsx")
    _desempenho_sheet(xl)
    src = CellMapAircraftSource(os.path.join(root, "data", "mappings", "dados_desempenho.json"))
    ac = src.load(xl)
    assert not src.warnings
    assert ac.wing.area == pytest.approx(0.9) and ac.wing.span == pytest.approx(2.4)
    assert ac.wing.CLa == pytest.approx(0.08 * 180 / math.pi)
    assert ac.surface("htail").eps0 == 1.5 and ac.tanks == []
    out = str(tmp_path / "out.xlsx")
    src.save(ac, out)
    same_spec(compile_aircraft(ac), compile_aircraft(src.load(out)))
    r = manager.run("takeoff", ac)
    assert r.status == "COMPLETED"


def test_table_catalog(tmp_path):
    from openpyxl import Workbook
    from aerosim.data.cellmap import TableCatalogSource
    wb = Workbook()
    m = wb.active
    m.title = "Motores_Data"
    m.append(["Nome", "Kv", "Rm (ohm)", "Imax (A)", "Massa", "Outra"])
    m.append(["M1", 600, 0.03, 40, 0.3, 1])
    p = wb.create_sheet("Propellers_coefficients")
    p.append(["titulo"])
    p.append(["Nome", "D", "passo", "p00"])
    p.append(["P1", 0.3, 0.15, 1.0])
    t = wb.create_sheet("Propellers_JCTCP")
    t.append(["helice", "J", "CT", "CP"])
    for j in (0.0, 0.3, 0.6):
        t.append(["P1", j, 0.1 - 0.15 * j, 0.04 - 0.03 * j])
    path = str(tmp_path / "pm.xlsx")
    wb.save(path)
    src = TableCatalogSource({"motors": {"sheet": "Motores_Data", "header_row": 1},
                              "propellers": {"sheet": "Propellers_coefficients", "header_row": 2},
                              "prop_table": {"sheet": "Propellers_JCTCP", "header_row": 1}})
    cat = src.load_catalog(path)
    assert cat.motors[0].Kv == 600 and cat.motors[0].Rm == 0.03
    assert cat.propellers[0].J == [0.0, 0.3, 0.6]
    assert any("Outra" in w for w in src.warnings) and any("p00" in w for w in src.warnings)
    _ = json
