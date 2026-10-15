"""GUI: monta a janela sem tela (offscreen), edita, arrasta e roda uma simulação pela aba."""
import os
import time

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
QtWidgets = pytest.importorskip("PyQt6.QtWidgets")


@pytest.fixture(scope="module")
def app():
    return QtWidgets.QApplication.instance() or QtWidgets.QApplication([])


def pump(app, t=0.2):
    end = time.time() + t
    while time.time() < end:
        app.processEvents()
        time.sleep(0.005)


def test_window_tabs_and_edit(app, manager):
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import AppState
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    w = MainWindow(st)
    w.show()
    pump(app)
    titles = [w.tabs.tabText(i) for i in range(w.tabs.count())]
    assert titles[-1] == "Montagem" and "Missão" in titles and "Desempenho" in titles
    assert "Decolagem" not in titles and "Polar" not in titles      # viraram modos/partes de outras abas
    cfg = w.config
    n0 = len(st.aircraft.propulsion)
    cfg._add_motor()
    cfg._add_tank()
    assert len(st.aircraft.propulsion) == n0 + 1 and len(st.aircraft.tanks) == 2
    cfg._rotate_wind(15)
    assert st.aircraft.environment.wind_from == 45
    from aerosim.model.geometry import aircraft_meshes
    cfg._add_boom()                                       # + Tubo, duas vezes: tubo duplo
    cfg._add_boom()
    b0, b1 = st.aircraft.booms
    assert b1.y == pytest.approx(b0.y + 0.3) and {"boom:0", "boom:1"} <= set(aircraft_meshes(st.aircraft))
    cfg._on_pick("boom:1")                                # o editor da direita é o do tubo escolhido
    assert cfg.form.obj is b1 and "pitch" in cfg.form.widgets
    cfg._on_pick("gear")                                  # tipo do trem também só pelo editor
    cfg.form.widgets["layout"].findChild(QtWidgets.QComboBox).setCurrentText("conventional")
    g = st.aircraft.gear
    assert g.layout == "conventional" and g.steer_x < g.main_x
    pump(app, 0.7)
    st.undo()                                             # Ctrl-Z volta o trem
    assert st.aircraft.gear.layout == "tricycle"
    st.redo()
    assert st.aircraft.gear.layout == "conventional"
    pump(app, 0.5)
    txt = cfg.summary.toPlainText()
    assert "margem" in txt and "Pista" in txt and "T/W" in txt
    cfg.restore()                                         # volta ao que foi aberto
    pump(app, 0.3)
    assert len(st.aircraft.propulsion) == n0 and len(st.aircraft.tanks) == 1


def test_simulation_tab_runs(app, manager):
    from aerosim.gui.state import AppState
    from aerosim.gui.tabs.simulation_tab import SimulationTab
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    tab = SimulationTab(st, manager.get("mission"))
    tab.inputs["mode"].setCurrentText("só decolagem")
    pump(app, 0.1)
    tab.reconfigure()
    assert "x_warning" in tab.inputs and "start" not in tab.inputs   # o formulário mudou com o modo
    tab.start()
    t0 = time.time()
    while tab._thread is not None and time.time() - t0 < 60:
        pump(app, 0.05)
    assert tab.result is not None and tab.result.status == "COMPLETED"
    assert tab.table.rowCount() > 5


def test_propulsion_poly_editor(app, manager, desempenho_xlsx, mapping_path):
    """importar a planilha traz os conjuntos de tração; o editor troca motor <-> polinômio"""
    from PyQt6 import QtCore
    from aerosim.data.cellmap import CellMapAircraftSource
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import AppState
    src = CellMapAircraftSource(mapping_path)
    ac = src.load(desempenho_xlsx)
    st = AppState(manager, ac)
    st.catalog = st.catalog.merge(src.catalog)
    w = MainWindow(st)
    cfg = w.config
    for i in range(cfg.comp_list.count()):
        if cfg.comp_list.item(i).data(QtCore.Qt.ItemDataRole.UserRole) == "prop:0":
            cfg.comp_list.setCurrentRow(i)
    pump(app)
    combos = cfg.editor_area.widget().findChildren(QtWidgets.QComboBox)
    model = next(c for c in combos if c.count() == 2 and "polinômio" in c.itemText(1))
    assert model.currentIndex() == 1
    sets = next(c for c in combos if c.findText("conjunto A") >= 0)
    sets.setCurrentText("conjunto A")
    assert st.aircraft.propulsion[0].thrust_poly.name == "conjunto A"
    model.setCurrentIndex(0)
    assert st.aircraft.propulsion[0].thrust_model == "motor"


def test_envelope_export_xlsx(app, manager, tmp_path, desempenho_xlsx, mapping_path):
    """o envelope sai em Excel (como o Envelope.xlsx do original): uma aba por tabela"""
    from openpyxl import load_workbook
    from aerosim.data.cellmap import CellMapAircraftSource
    from aerosim.gui.tabs.simulation_tab import export_result
    ac = CellMapAircraftSource(mapping_path).load(desempenho_xlsx)
    res = manager.run("envelope", ac, {"method": "planilha", "table": False})
    out = str(tmp_path / "Envelope.xlsx")
    export_result(res, out)
    wb = load_workbook(out)
    sh = wb["envelope_planilha"]
    head = [c.value for c in sh[1]]
    assert head[:7] == ["h", "rho", "Vmin", "Vmax", "Ve", "Vman", "Vcru"]
    assert sh.max_row == len(res.series["envelope_planilha"]["h"]) + 1
    assert "resumo" in wb.sheetnames


def test_stability_and_aero_tabs(app, manager):
    from aerosim.gui.state import AppState
    from aerosim.gui.tabs.simulation_tab import SimulationTab
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    for name, over in (("stability", {"size_tail": False}), ("aerodynamics", {"use_lltdrek": False})):
        tab = SimulationTab(st, manager.get(name))
        for k, v in over.items():
            tab.inputs[k].setChecked(v)
        tab.start()
        t0 = time.time()
        while tab._thread is not None and time.time() - t0 < 120:
            pump(app, 0.05)
        assert tab.result is not None, tab.msg.text()
        assert tab.table.rowCount() > 3 and tab.diag.rowCount() >= 3
    # a aba se refaz quando a configuração muda (tirar a EH tira a área de EH)
    tab = SimulationTab(st, manager.get("stability"))
    assert "size_tail" in tab.inputs
    st.aircraft.surfaces = [s for s in st.aircraft.surfaces if s.role != "htail"]
    st.notify("geometry")
    tab.reconfigure()
    assert "size_tail" not in tab.inputs


def test_undo_restores_tab_values(app, manager):
    """Ctrl+Z também volta o que se escolheu nos formulários e tabelas das abas"""
    from aerosim.gui.state import AppState
    from aerosim.gui.tabs.simulation_tab import SimulationTab
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    tab = SimulationTab(st, manager.get("mission"))
    tab.inputs["mode"].setCurrentText("só decolagem")
    pump(app, 0.1)
    assert "start" not in tab.inputs
    assert st.undo()
    pump(app, 0.1)
    assert tab.inputs["mode"].currentText() == "missão" and "start" in tab.inputs
    assert st.redo()
    pump(app, 0.1)
    assert tab.inputs["mode"].currentText() == "só decolagem"


def test_sprint_stages_from_scratch_to_assembly(app, manager, tmp_path, monkeypatch):
    """Perfil → Asa → Empenagem → Motores: cada etapa guarda e a última manda o avião para a Montagem"""
    monkeypatch.setenv("AEROSIM_PROJETOS", str(tmp_path))
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import AppState
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    w = MainWindow(st)
    w.show()
    titles = [w.tabs.tabText(i) for i in range(w.tabs.count())]
    assert titles.index("Aerodinâmica") < titles.index("Estabilidade") < titles.index("Desempenho") < titles.index("Missão")
    aero = w.tabs.widget(titles.index("Aerodinâmica"))
    assert [aero.tabText(i) for i in range(aero.count())] == ["Perfil", "Asa", "Avião montado"]
    w.tabs.setCurrentWidget(aero)
    perfil = aero.widget(0)
    pump(app, 0.6)
    assert perfil.checks.table.rowCount() > 3              # a polar de cada candidato
    perfil._keep()                                         # "Usar na asa" vai para a etapa Asa
    assert aero.currentIndex() == 1
    asa = aero.widget(1)
    asa.form.widgets["area"].setValue(1.0)
    asa.form.widgets["area"].valueChanged.emit(1.0)
    pump(app, 0.6)
    assert st.design.wing_draft.area == 1.0 and asa.checks.diag.rowCount() >= 3
    asa._keep()
    assert [x.name for x in st.design.wings] == ["asa 1"]
    est = w.tabs.widget(titles.index("Estabilidade"))
    w.tabs.setCurrentWidget(est)
    pump(app, 0.5)
    est.widget(0)._keep()
    assert st.design.sets[0].wing.name == "asa 1"
    des = w.tabs.widget(titles.index("Desempenho"))
    w.tabs.setCurrentWidget(des)
    pump(app, 1.0)
    motores = des.widget(0)
    assert motores.ac is not None and motores.checks.diag.rowCount() >= 3
    assert [des.tabText(i) for i in range(des.count())] == ["Motores e montagem", "Estudo e ponto de projeto",
                                                            "Avião montado"]
    estudo = des.widget(1)                                 # retroalimentação: ponto de projeto acoplado
    st.design.study.n, st.design.study.places = 2, "A=600"
    st.design.requirements.V_stall_max = 13.0            # estol no MTOW que este avião cumpre (restrição ativa)
    st.design.requirements.takeoff_run = 50.0            # cenário com pista longa (o padrão SAE é 6 m)
    estudo.start()
    t0 = time.time()
    while estudo._thread is not None and time.time() - t0 < 120:
        pump(app, 0.1)
    assert len(estudo.rows) == 2 and estudo.table.rowCount() == 2
    assert len(st.design.studies) == 1 and len(st.design.studies[0].rows) == 2     # dados de simulação no projeto
    from aerosim.gui.design.pages import StudyPage
    reaberto = StudyPage(st)                               # reabrir: o estudo guardado volta sem rodar
    assert reaberto.table.rowCount() == 2 and "_cand" not in reaberto.rows[0]
    reaberto._show_row(0)                                  # refaz o ponto de projeto só quando precisa
    assert "_cand" in reaberto.rows[0]
    n = len(st.fleet)
    motores.send()
    assert len(st.fleet) == n + 1 and w.tabs.currentWidget() is w.config
    assert st.aircraft.name.startswith("avião 1") and st.path and os.path.exists(st.path)
    pump(app, 1.2)
    assert os.path.exists(os.path.join(st.design_dir, "sprint.json"))   # o projeto da sprint se grava sozinho
    asas = [n for n in os.listdir(st.design_dir) if n != "sprint.json"]
    assert asas and all(n.endswith(".json") for n in asas)              # um arquivo por asa
    import aerosim.design.flight as fl
    real = fl.tune_takeoff
    fl.tune_takeoff = lambda ac, backend: (False, "não decola na pista (teste)")
    try:
        motores.send(go="desempenho")                     # não decola: vai assim mesmo, com o aviso
    finally:
        fl.tune_takeoff = real
    assert len(st.fleet) == n + 2 and "⚠" in motores.msg.text()
    perf = w.area_tabs["desempenho"]
    assert w.tabs.currentWidget() is perf and perf.currentIndex() == perf.count() - 1


def test_open_project_resumes_where_it_stopped(app, manager, tmp_path):
    """abrir o .json do projeto volta de onde parou: Montagem (aviões abertos) e aba; sem avião guardado, vazia"""
    from aerosim.data.projects import JsonProjectStore
    from aerosim.design.project import DesignProject
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import AppState
    from aerosim.model.templates import empty, is_empty, twin_boom
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    w = MainWindow(st)
    w.show()
    pump(app)
    assert is_empty(st.aircraft)
    w.tabs.setCurrentIndex(w.tabs.count() - 1)
    w.config._add_boom()                                  # vazio → as peças entram uma a uma
    assert len(st.aircraft.booms) == 1
    p = DesignProject()
    p.fleet, p.current, p.ui = [empty(), twin_boom()], 1, {"tab": "Desempenho", "sub": 1}
    path = str(tmp_path / "meu_projeto.json")
    p.save(path)
    before = open(path, "rb").read()
    w.open_path(path, quiet=True)
    pump(app)
    folder = str(tmp_path / "meu_projeto")                # o projeto antigo passa a ser gravado numa pasta ao lado
    assert st.names() == [a.name for a in p.fleet] and st.current == 1 and st.design_dir == folder
    assert w.ui_state() == {"tab": "Desempenho", "sub": 1}
    st.save_design()                                      # grava de volta o que está na tela
    from aerosim.design import store
    assert store.load(folder).current == 1 and open(path, "rb").read() == before
    w.open_path(folder, quiet=True)                       # e a pasta abre do mesmo jeito
    assert st.current == 1 and st.design_dir == folder
    empty_path = str(tmp_path / "sem_avioes.json")
    DesignProject().save(empty_path)
    w.open_path(empty_path, quiet=True)
    assert len(st.fleet) == 1 and is_empty(st.aircraft)
    w.close()


def test_control_mix_row(app, manager):
    """a mistura de um comando se edita na Montagem; escolher um tipo conhecido volta a mistura dele"""
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import AppState
    from aerosim.model.templates import conventional
    st = AppState(manager, conventional())
    w = MainWindow(st)
    cfg = w.config
    ail = st.aircraft.wing.controls[0]
    from aerosim.gui.widgets.param_slider import DataclassForm
    row = cfg._mix_row(ail, DataclassForm(ail))
    spins = row.findChildren(QtWidgets.QDoubleSpinBox)
    assert [s.value() for s in spins] == [0.0, 1.0, 0.0]
    spins[0].setValue(1.0)                                # aileron + arfagem = elevon
    assert ail.mix == [1.0, 1.0, 0.0] and ail.role == "elevon"
    w.close()


def test_constraints_dialog(app, manager, tmp_path):
    """restrições ligam/desligam e mudam as contas; variáveis escrevem no rascunho da etapa"""
    from aerosim.data.projects import JsonProjectStore
    from aerosim.gui.design.pages import RequirementsDialog
    from aerosim.gui.state import AppState
    from aerosim.model.templates import empty
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    dlg = RequirementsDialog(st)
    req = st.design.requirements
    t = dlg.constraints.table
    row = next(i for i, c in enumerate(req.constraints) if c.key == "s_to")
    t.cellWidget(row, 3).setValue(3.0)                       # decolar em 3 m
    assert req.takeoff_run == 3.0
    dlg.constraints.table.cellWidget(row, 0).setChecked(False)
    assert req.takeoff_run > 100                             # desligada: sem limite
    vt = dlg.variables.table
    r = next(i for i, v in enumerate(st.design.variables) if v.key == "wing.span")
    vt.cellWidget(r, 2).setValue(1.8)
    vt.cellWidget(r, 1).setCurrentText("variável")
    assert st.design.wing_draft.span == 1.8 and st.design.variables[r].mode == "variável"
    dlg.close()


def test_mdo_page_realtime_history(app, manager, tmp_path):
    """mini-MDO: o otimizador do núcleo avisa a cada iteração; slider do histórico; clique; levar para as etapas"""
    from aerosim.data.projects import JsonProjectStore
    from aerosim.gui.design.mdo_page import MdoPage
    from aerosim.gui.state import AppState
    from aerosim.model.templates import empty
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    d = st.design
    d.airfoils.append(__import__("aerosim.design.airfoils", fromlist=["entry"]).entry(
        __import__("aerosim.design.airfoils", fromlist=["from_naca"]).from_naca("NACA 0012")))
    for v in d.variables:
        if v.key in ("wing.span", "wing.area", "wing.taper"):
            v.mode = "variável"
        if v.key == "wing.airfoil_tip":
            v.mode, v.options = "variável", ["NACA 4412", "NACA 0012"]
    d.mdo.particles, d.mdo.iterations = 8, 6
    page = MdoPage(st)
    seen = []
    redraw = page.redraw
    page.redraw = lambda: (seen.append(page.slider.maximum()), redraw())
    page.start()
    t0 = time.time()
    while page._thread is not None and time.time() - t0 < 60:
        pump(app, 0.05)
    assert len(page.rows) == 48 and len(page.hist) == 6
    assert len(seen) >= 6 and len(set(seen)) >= 3             # redesenhou a cada iteração (tempo real)
    page.slider.setValue(0)
    assert page.iter_lbl.text().startswith("1 /")
    assert {r["airfoil_tip"] for r in page.rows} <= {"NACA 4412", "NACA 0012", "NACA 2412"}
    page.select(page.hist[-1])
    assert page.table.rowCount() > 10
    page.apply()
    assert d.wing_draft.span == pytest.approx(round(page.hist[-1]["span"], 3))
    assert len(d.mdo_runs) == 1
    again = MdoPage(st)                                        # reabrir: o resultado guardado volta
    assert len(again.rows) == 48


def test_mdo_bounds_sheet(app, manager, tmp_path):
    """sem duplicata: as variáveis do Mini-MDO são as de Condições e restrições; restrições viram limites; limite
    inválido não roda"""
    from aerosim.data.projects import JsonProjectStore
    from aerosim.design import registry as rg
    from aerosim.gui.design.mdo_page import MdoPage
    from aerosim.gui.design.pages import RequirementsDialog
    from aerosim.gui.state import AppState
    from aerosim.model.templates import empty
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    d = st.design
    page = MdoPage(st)
    assert not hasattr(page, "vtable")                       # a planilha mora no diálogo
    dlg = RequirementsDialog(st)
    d.requirements.constraints = [rg.Constraint("b", "≤", 2.2), rg.Constraint("S", "≥", 0.6),
                                  rg.Constraint("n_motor", "≤", 2)]
    dlg.variables.bounds_b.click()
    get = lambda k: next(v for v in d.variables if v.key == k)
    assert get("wing.span").hi == 2.2 and get("wing.area").lo == 0.6 and get("power.motors").hi == 2
    span = get("wing.span")
    span.mode, span.lo, span.hi = "variável", 2.0, 1.0
    page.start()
    assert page._thread is None and "limites inválidos" in page.msg.text()



def test_list_picker_search_and_select_all(app, manager, tmp_path):
    """perfil, motor, hélice, arranjo: escolher numa lista com busca; marcar todos marca só o que a busca mostra"""
    from PyQt6 import QtCore
    from aerosim.data.projects import JsonProjectStore
    from aerosim.design import mdo
    from aerosim.gui.design.pages import ListPicker, VariablesTable
    from aerosim.gui.state import AppState
    from aerosim.model.templates import empty
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    items = mdo.available(st.design, "wing.airfoil")
    assert any("S1223" in n for n in items) and any("NACA 0012" in n for n in items)   # a pasta de perfis
    dlg = ListPicker("perfil", items, [])
    assert dlg.selected() == []                                  # tudo marcado = todos
    dlg.none_b.click()
    dlg.search.setText("naca")
    dlg.all_b.click()                                            # só os NACA
    sel = dlg.selected()
    assert sel and all("naca" in s.lower() for s in sel)
    assert set(mdo.available(st.design, "power.motor")) == {m.name for m in st.catalog.motors}
    vt = VariablesTable(st)
    r = next(i for i, v in enumerate(st.design.variables) if v.key == "power.propeller")
    assert vt.table.cellWidget(r, 5).text().startswith("todos (")


def test_repo_sample_project_saves_outside_data(app, manager, tmp_path):
    """o exemplo de data/sprint abre, mas grava na pasta da sprint do usuário (nada em data/)"""
    from aerosim.data.projects import JsonProjectStore
    from aerosim.gui.main_window import MainWindow
    from aerosim.gui.state import ROOT, AppState
    from aerosim.model.templates import empty
    sample = os.path.join(ROOT, "data", "sprint", "projeto_s1223.json")
    before = sorted(os.listdir(os.path.dirname(sample)))
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    w = MainWindow(st)
    w.open_path(sample, quiet=True)
    st.save_design()
    assert st.design_dir == os.path.join(str(tmp_path), "sprint", "projeto_s1223")
    assert "asa S1223.json" in os.listdir(st.design_dir)
    assert sorted(os.listdir(os.path.dirname(sample))) == before
    w.close()


def test_save_checklist_dialog_and_refusal(app, manager, tmp_path):
    """antes de salvar: a lista (caixas só de leitura) e 'salvar mesmo assim?'; recusar não salva nem abre"""
    from PyQt6 import QtCore
    from aerosim.data.projects import JsonProjectStore
    from aerosim.design import checklist as ck
    from aerosim.gui.design import pages
    from aerosim.gui.state import AppState
    from aerosim.gui.widgets import save_check
    from aerosim.model.templates import conventional, empty
    ac = conventional()
    ac.payload.enabled, ac.payload.mass = True, 0.0
    checks = ck.aircraft_checklist(ac, None)
    dlg = save_check.SaveCheckDialog(ac, checks)
    g = dlg.tree.topLevelItem(0)
    states = {g.child(i).text(0)[2:]: g.child(i).checkState(0) for i in range(g.childCount())}
    assert states["LAPES: massa"] == QtCore.Qt.CheckState.Unchecked and states["asa"] == QtCore.Qt.CheckState.Checked
    assert not (g.child(0).flags() & QtCore.Qt.ItemFlag.ItemIsUserCheckable)       # só leitura
    assert dlg.only.isChecked()                                                    # abre mostrando as pendências
    hidden = [g.child(i).isHidden() for i in range(g.childCount())]
    assert any(hidden) and not all(hidden)
    dlg.save_b.click()
    assert dlg.result() == 1
    st = AppState(manager, empty(), projects=JsonProjectStore(str(tmp_path)))
    n = len(st.fleet)
    real = save_check.confirm_save
    save_check.confirm_save = lambda parent, a, project=None: False
    try:
        w = QtWidgets.QWidget()
        assert pages.send_to_assembly(st, ac, w) is False
    finally:
        save_check.confirm_save = real
    assert len(st.fleet) == n and not st.projects.exists(ac.name)              # recusou: nada salvo nem aberto
